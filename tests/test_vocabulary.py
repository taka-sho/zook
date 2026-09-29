"""Type lookup forgiveness and the hints an unresolved type gets - what lets
an AI fix a wrong `type` from the Warning alone."""

import pytest

from zook.layout import build_layout, icon_resolution_warnings
from zook.model import parse_diagram
from zook.registry import BUILTIN_PROVIDERS, load_registries, normalize_type
from zook.validate import validate

REGISTRY = load_registries()


def _warnings(elements):
    doc = {"version": "1.0", "canvas": {"aspectRatio": "16:9"}, "elements": elements}
    validate(doc)
    diagram = parse_diagram(doc)
    return icon_resolution_warnings(build_layout(diagram, REGISTRY), REGISTRY)


@pytest.mark.parametrize("written", ["API Gateway", "api-gateway", "api_gateway", "apigateway", "APIGateway"])
def test_spelling_variants_resolve(written):
    assert REGISTRY.resolve_icon(written, "aws").name == "APIGateway"


def test_no_two_entries_share_a_lookup_key():
    # normalize_type() folds case and punctuation: two different entries
    # of one provider must never collapse onto the same key.
    import yaml
    from importlib import resources

    for provider in BUILTIN_PROVIDERS:
        raw = yaml.safe_load(
            resources.files(f"zook.data.icons.{provider}").joinpath(f"registry.{provider}.yaml").read_text(encoding="utf-8")
        )
        for section in ("icons", "groups"):
            owner = {}
            for name, spec in raw.get(section, {}).items():
                for key in {normalize_type(n) for n in [name, *spec.get("aliases", [])]}:
                    assert owner.setdefault(key, name) == name, f"{provider} {section}: {key!r} is both {owner[key]} and {name}"


def test_actors_resolve_for_every_provider():
    for provider in BUILTIN_PROVIDERS:
        assert REGISTRY.resolve_icon("User", provider).name == "User"
    assert _warnings([{"kind": "node", "id": "u", "type": "Admin", "provider": "gcp"}]) == []


def test_a_typo_gets_the_close_match():
    (message,) = _warnings([{"kind": "node", "id": "a", "type": "Lamda"}])
    assert message.endswith("did you mean 'Lambda'?")


def test_another_providers_type_says_which_provider():
    (message,) = _warnings([{"kind": "node", "id": "a", "type": "CloudRun"}])
    assert "set `provider: gcp`" in message


def test_an_unknown_type_points_at_the_vocabulary_and_plain_shapes():
    (message,) = _warnings([{"kind": "node", "id": "a", "type": "QuantumFlux"}])
    assert "zook icons list" in message and "shape: rect" in message


def test_an_unknown_container_type_is_a_warning():
    (message,) = _warnings([{"kind": "container", "id": "n", "type": "vcp", "children": [{"kind": "node", "id": "a", "type": "EC2"}]}])
    assert message.startswith("unknown container type 'vcp'") and "did you mean 'vpc'?" in message


def test_every_builtin_icon_and_badge_file_exists():
    for provider, registry in REGISTRY.registries.items():
        for entry in registry.icons.values():
            assert entry.file.is_file(), f"{provider}: {entry.name} -> {entry.file}"
        for group in registry.groups.values():
            assert group.icon is None or group.icon.is_file(), f"{provider}: {group.name}"


@pytest.mark.parametrize(
    "type_, expected",
    [("WAF", "WAF"), ("Secrets Manager", "SecretsManager"), ("CloudWatch", "CloudWatch"), ("IGW", "InternetGateway"),
     ("Step Functions", "StepFunctions"), ("Kinesis", "Kinesis"), ("OpenSearch", "OpenSearch"), ("Bedrock", "Bedrock")],
)
def test_common_aws_services_beyond_tier_1(type_, expected):
    assert REGISTRY.resolve_icon(type_, "aws").name == expected


@pytest.mark.parametrize("type_", ["Server", "Internet", "OnPremises", "SaaS", "Mobile", "Database"])
def test_generic_icons_serve_every_provider(type_):
    for provider in BUILTIN_PROVIDERS:
        assert REGISTRY.resolve_icon(type_, provider) is not None


@pytest.mark.parametrize("group", ["publicSubnet", "private-subnet", "securityGroup", "asg", "corporateDataCenter", "availability zone"])
def test_aws_group_vocabulary(group):
    assert REGISTRY.resolve_group(group, "aws") is not None


def _custom_icon_setup(tmp_path, icon_name, icon_bytes):
    (tmp_path / icon_name).write_bytes(icon_bytes)
    registry = tmp_path / "reg.yaml"
    registry.write_text(
        f'registryVersion: "1.0"\nprovider: aws\nicons:\n  MySvc: {{ file: "{icon_name}" }}\n', encoding="utf-8"
    )
    diagram = tmp_path / "d.yaml"
    diagram.write_text(
        'version: "1.0"\ncanvas: {aspectRatio: "16:9"}\nelements:\n  - {kind: node, id: a, type: MySvc}\n',
        encoding="utf-8",
    )
    return diagram, registry


SVG = b'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 10 10"><rect width="10" height="10" fill="#f60"/></svg>'


def test_an_svg_icon_is_rasterized_for_every_output(tmp_path):
    import json

    from click.testing import CliRunner
    from pptx import Presentation

    from zook.cli import main

    diagram, registry = _custom_icon_setup(tmp_path, "my.svg", SVG)
    runner = CliRunner()
    for args in (["build", str(diagram), "-o", str(tmp_path / "o.pptx")],
                 ["preview", str(diagram), "-o", str(tmp_path / "o.png")],
                 ["export-drawio", str(diagram), "-o", str(tmp_path / "o.drawio")]):
        result = runner.invoke(main, [*args, "--registry", str(registry), "--format", "json"])
        assert result.exit_code == 0, result.output
        assert json.loads(result.stdout)["status"] == "ok"
    def walk(shapes):
        for shape in shapes:
            yield shape
            if shape.shape_type == 6:
                yield from walk(shape.shapes)

    pictures = [s for s in walk(Presentation(str(tmp_path / "o.pptx")).slides[0].shapes) if s.shape_type == 13]
    assert pictures and pictures[0].image.content_type == "image/png"
    assert "image=data:image/png," in (tmp_path / "o.drawio").read_text(encoding="utf-8")


def test_an_unreadable_icon_file_is_a_warning_not_a_crash(tmp_path):
    import json

    from click.testing import CliRunner

    from zook.cli import main

    diagram, registry = _custom_icon_setup(tmp_path, "my.png", b"not an image")
    result = CliRunner().invoke(main, ["build", str(diagram), "-o", str(tmp_path / "o.pptx"), "--registry", str(registry), "--format", "json"])
    assert result.exit_code == 0, result.output
    (message,) = json.loads(result.stdout)["warnings"]
    assert "is not an image zook can read" in message and "placeholder" in message


def test_placeholder_icons_are_distinguishable_within_a_provider():
    # The generated placeholders used to abbreviate to the first four letters,
    # so CloudRun/CloudSQL/CloudDNS/... all read "CLOU" (some byte-identical).
    import hashlib

    for provider, registry in REGISTRY.registries.items():
        seen = {}
        for entry in {id(e): e for e in registry.icons.values()}.values():
            if entry.category == "General":
                continue  # the actors share one person glyph on purpose
            digest = hashlib.sha256(entry.file.read_bytes()).hexdigest()
            assert seen.setdefault(digest, entry.name) == entry.name, f"{provider}: {entry.name} looks like {seen[digest]}"
