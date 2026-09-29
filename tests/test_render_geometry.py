"""What the generated .pptx actually draws, checked from its stored geometry.

A connector's route isn't stored as points: it is a preset shape
(bentConnector3 for an elbow) whose path is defined in the shape's own
coordinate frame and then flipped and rotated by its xfrm. These tests
reconstruct the drawn path from that XML - exactly as PowerPoint does - and
compare it with the route layout planned (and every crossing check and the
preview assumed). They also cover the text boxes, fit-to-slide, and the
top-level arrangement.
"""

import math

import pytest
from pptx.util import Emu, Pt

from zook.layout import (
    build_layout,
    canvas_warnings,
    diagram_warnings,
    fit_transform,
    iter_boxes,
    link_aliasing_warnings,
    link_label_rect,
    link_label_rect_for,
    _rects_overlap,
    link_render_plan,
    node_label_rect,
)
from zook.model import parse_diagram
from zook.registry import load_registries
from zook.render import LOGICAL_TO_EMU, render
from zook.validate import validate

REGISTRY = load_registries()


def _doc(elements, links=None, **canvas):
    doc = {"version": "1.0", "canvas": {"aspectRatio": "16:9", **canvas}, "elements": elements}
    if links is not None:
        doc["links"] = links
    validate(doc)
    return doc


def _node(eid, x=None, y=None, **extra):
    node = {"kind": "node", "id": eid, "type": "EC2", **extra}
    if x is not None:
        node.update(x=x, y=y)
    return node


def _build(doc):
    diagram = parse_diagram(doc)
    root = build_layout(diagram, REGISTRY)
    return diagram, root, render(diagram, root, REGISTRY)


def _all_shapes(shapes):
    """Every shape on the slide, inside groups too (a node's icon and label
    are grouped; containers are groups)."""
    for shape in shapes:
        yield shape
        if shape.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP
            yield from _all_shapes(shape.shapes)


def _connectors(prs):
    return [s for s in prs.slides[0].shapes if s.shape_type is not None and s._element.tag.endswith("cxnSp")]


def _drawn_path(conn) -> list[tuple[float, float]]:
    """The polyline a viewer draws for a straight/bentConnector3 cxnSp, in
    logical units: preset path in the local frame -> flipH/flipV -> rot about
    the centre -> offset."""
    xfrm = conn._element.spPr.xfrm
    ox, oy = xfrm.off.x, xfrm.off.y
    w, h = xfrm.ext.cx, xfrm.ext.cy
    prst = conn._element.spPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}prstGeom").get("prst")
    local = [(0, 0), (w, h)] if prst in ("line", "straightConnector1") else [(0, 0), (w / 2, 0), (w / 2, h), (w, h)]
    rot = int(xfrm.get("rot", "0")) / 60000
    flip_h, flip_v = xfrm.get("flipH") == "1", xfrm.get("flipV") == "1"
    cx, cy = w / 2, h / 2
    out = []
    for x, y in local:
        x, y = x - cx, y - cy
        if flip_h:
            x = -x
        if flip_v:
            y = -y
        r = math.radians(rot)
        x, y = x * math.cos(r) - y * math.sin(r), x * math.sin(r) + y * math.cos(r)
        out.append(((ox + cx + x) / LOGICAL_TO_EMU, (oy + cy + y) / LOGICAL_TO_EMU))
    return out


def _close(path_a, path_b, tol=1.0):
    return len(path_a) == len(path_b) and all(math.dist(p, q) <= tol for p, q in zip(path_a, path_b))


@pytest.mark.parametrize(
    "a_xy,b_xy,from_side,to_side",
    [
        ((100, 100), (400, 400), "bottom", "top"),  # down-right
        ((400, 100), (100, 400), "bottom", "top"),  # down-left
        ((400, 400), (100, 100), "top", "bottom"),  # up-left
        ((100, 400), (400, 100), "top", "bottom"),  # up-right
        ((100, 100), (400, 300), "right", "left"),  # horizontal elbows, both directions
        ((400, 300), (100, 100), "left", "right"),
    ],
)
def test_elbow_connector_draws_the_planned_route(a_xy, b_xy, from_side, to_side):
    # A vertical elbow used to be written unrotated, which any viewer that draws
    # the stored geometry renders H-V-H instead of the planned V-H-V.
    doc = _doc(
        [_node("a", *a_xy, style={"labelPosition": "none"}), _node("b", *b_xy, style={"labelPosition": "none"})],
        [{"from": "a", "to": "b", "fromSide": from_side, "toSide": to_side}],
    )
    diagram, root, prs = _build(doc)
    boxes = {b.element.id: b for b in iter_boxes(root)}
    _, _, style, planned = link_render_plan(boxes["a"], boxes["b"], diagram.links[0])
    assert style == "elbow"
    (conn,) = _connectors(prs)
    assert _close(_drawn_path(conn), planned), (_drawn_path(conn), planned)


def test_same_side_link_is_drawn_as_the_planned_u_route():
    doc = _doc(
        [_node("a", 100, 200, style={"labelPosition": "none"}), _node("b", 400, 260, style={"labelPosition": "none"})],
        [{"from": "a", "to": "b", "fromSide": "top", "toSide": "top"}],
    )
    diagram, root, prs = _build(doc)
    boxes = {b.element.id: b for b in iter_boxes(root)}
    _, _, style, planned = link_render_plan(boxes["a"], boxes["b"], diagram.links[0])
    assert style == "polyline"
    assert planned[1][1] < min(planned[0][1], planned[-1][1])  # goes up and over, not through a
    drawn = [pt for conn in _connectors(prs) for pt in _drawn_path(conn)]
    # consecutive segments share their joints; dedupe to the polyline itself
    polyline = [drawn[0]] + [p for i, p in enumerate(drawn[1:], 1) if not _close([p], [drawn[i - 1]])]
    assert _close(polyline, planned)


def test_node_label_box_is_sized_for_its_wrapped_text():
    doc = _doc([_node("a", 0, 0, label="アプリケーションロードバランサー\nprimary")])
    diagram, root, prs = _build(doc)
    box = next(b for b in iter_boxes(root) if b.element.id == "a")
    assert len(box.label_lines) >= 3  # the long line wraps, plus the explicit break
    assert box.footprint_h == pytest.approx(box.height + 4 + box.label_h)
    textbox = next(s for s in _all_shapes(prs.slides[0].shapes) if s.name == "a label")
    x, y, w, h = node_label_rect(box)
    assert textbox.height == Emu(round(h * LOGICAL_TO_EMU))
    assert textbox.text_frame.margin_left == 0  # measured against the full box width


def test_six_japanese_characters_fit_on_one_line():
    doc = _doc([_node("a", 0, 0, label="データベース")])
    _, root, _ = _build(doc)
    box = next(b for b in iter_boxes(root) if b.element.id == "a")
    assert box.label_lines == ["データベース"]


def test_long_label_is_detected_when_it_would_collide():
    # Two neighbours whose wrapped labels now take three lines each: the row
    # below must be pushed down by that, not overlap it.
    doc = _doc([
        {"kind": "container", "id": "g", "type": "group", "layout": {"direction": "vertical", "gap": 4},
         "children": [_node("a", label="a very long label that wraps onto several lines"), _node("b")]},
    ])
    diagram, root, _ = _build(doc)
    assert diagram_warnings(diagram, root, REGISTRY) == []


def test_short_labeled_link_keeps_its_arrowhead_visible():
    # Adjacent icons at the default gap: the centred white label box used to
    # cover the whole line and its arrowhead.
    doc = _doc([
        {"kind": "container", "id": "g", "type": "group", "layout": {"direction": "horizontal"},
         "children": [_node("a"), _node("b")]},
    ], [{"from": "a", "to": "b", "label": "SQL (3306)"}])
    diagram, root, _ = _build(doc)
    boxes = {b.element.id: b for b in iter_boxes(root)}
    _, _, _, path = link_render_plan(boxes["a"], boxes["b"], diagram.links[0])
    x, y, w, h = link_label_rect(path, "SQL (3306)", 8)
    end_x, end_y = path[-1]
    assert not (x <= end_x <= x + w and y <= end_y <= y + h)
    assert y + h < end_y  # moved above the horizontal line


@pytest.mark.parametrize("direction", ["horizontal", "vertical"])
def test_a_label_moved_beside_a_short_link_stays_off_both_icons(direction):
    # Wider than the gap between the icons: just beside the line it would sit
    # on top of both of them (opaque white box), with no warning.
    doc = _doc([
        {"kind": "container", "id": "g", "type": "group", "layout": {"direction": direction},
         "children": [_node("a"), _node("b")]},
    ], [{"from": "a", "to": "b", "label": "HTTPS + JWT authorization header"}])
    diagram, root, _ = _build(doc)
    boxes = {b.element.id: b for b in iter_boxes(root)}
    _, _, _, path = link_render_plan(boxes["a"], boxes["b"], diagram.links[0])
    label = link_label_rect_for(boxes["a"], boxes["b"], path, diagram.links[0])
    for node in ("a", "b"):
        icon = (boxes[node].abs_x, boxes[node].abs_y, boxes[node].width, boxes[node].height)
        assert not _rects_overlap(label, icon)
    assert diagram_warnings(diagram, root, REGISTRY) == []


def test_a_moved_label_stays_next_to_its_own_link():
    # Short link from a box into a diamond: the nearest free spot is above
    # the gap between them, not across the diamond beside the next link.
    doc = _doc([
        {"kind": "container", "id": "g", "type": "group", "layout": {"direction": "horizontal"},
         "children": [_node("a", label="Client", style={"shape": "rounded"}),
                      _node("b", label="Gateway", style={"shape": "diamond"}),
                      _node("c", label="Lambda", style={"shape": "rect"})]},
    ], [{"from": "a", "to": "b", "label": "HTTPS"}, {"from": "b", "to": "c"}])
    diagram, root, _ = _build(doc)
    boxes = {b.element.id: b for b in iter_boxes(root)}
    _, _, _, path = link_render_plan(boxes["a"], boxes["b"], diagram.links[0])
    x, _, w, _ = link_label_rect_for(boxes["a"], boxes["b"], path, diagram.links[0])
    a, b = boxes["a"], boxes["b"]
    assert a.abs_x + a.width - 1 <= x + w / 2 <= b.abs_x + 1
    assert diagram_warnings(diagram, root, REGISTRY) == []


def test_actor_next_to_a_wide_container_stays_on_the_slide():
    # With a uniform grid cell the actor's cell was as wide as the VPC, so the
    # VPC landed past the right edge of the slide.
    children = [_node(f"n{i}") for i in range(6)]
    doc = _doc([
        _node("user", type="User"),
        {"kind": "container", "id": "vpc", "type": "vpc", "label": "VPC",
         "layout": {"direction": "horizontal"}, "children": children},
    ])
    diagram, root, _ = _build(doc)
    assert canvas_warnings(diagram, root) == []
    assert fit_transform(diagram, root).is_identity
    boxes = {b.element.id: b for b in iter_boxes(root)}
    assert boxes["vpc"].abs_x < boxes["user"].abs_x + boxes["user"].footprint_w + 100


def test_canvas_layout_arranges_the_top_level():
    doc = _doc([_node("a"), _node("b"), _node("c")], layout={"direction": "vertical", "gap": 10})
    _, root, _ = _build(doc)
    boxes = {b.element.id: b for b in iter_boxes(root)}
    assert boxes["a"].abs_x == boxes["b"].abs_x == boxes["c"].abs_x
    assert boxes["a"].abs_y < boxes["b"].abs_y < boxes["c"].abs_y


def _too_big():
    return [_node(f"n{i}", i * 190, 0) for i in range(8)]  # ~1420 wide on a 1280 slide


def test_a_diagram_that_does_not_fit_is_shrunk_onto_the_slide():
    diagram, root, prs = _build(_doc(_too_big()))
    transform = fit_transform(diagram, root)
    assert 0.7 < transform.scale < 1.0
    assert canvas_warnings(diagram, root) == []  # a mild shrink isn't reported
    slide_w, slide_h = prs.slide_width, prs.slide_height
    for shape in prs.slides[0].shapes:
        assert 0 <= shape.left and shape.left + shape.width <= slide_w
        assert 0 <= shape.top and shape.top + shape.height <= slide_h


def test_a_heavy_shrink_is_reported():
    doc = _doc([_node(f"n{i}", i * 300, 0) for i in range(10)])  # ~2900 wide -> well under 70%
    diagram, root, _ = _build(doc)
    (message,) = canvas_warnings(diagram, root)
    assert "scaled to" in message


def test_a_heavy_shrink_names_a_stray_explicit_coordinate():
    # One typo'd x on an otherwise small diagram: the Warning must point at it,
    # not suggest splitting the slide.
    doc = _doc([_node("a", 100, 100), _node("b", 12000, 100)])
    diagram, root, _ = _build(doc)
    (message,) = canvas_warnings(diagram, root)
    assert "scaled to" in message and "'b' (x=12000" in message and "fewer elements" not in message


def test_fit_none_reports_what_falls_off_the_slide():
    diagram, root, _ = _build(_doc(_too_big(), fit="none"))
    assert fit_transform(diagram, root).is_identity
    assert any("outside the canvas" in m for m in canvas_warnings(diagram, root))


def test_a_diagram_that_fits_is_not_moved():
    doc = _doc([_node("a", 100, 100)])
    diagram, root, prs = _build(doc)
    assert fit_transform(diagram, root).is_identity
    pic = next(s for s in _all_shapes(prs.slides[0].shapes) if s.name == "a icon")
    assert (pic.left, pic.top) == (Emu(100 * LOGICAL_TO_EMU), Emu(100 * LOGICAL_TO_EMU))


def test_fan_out_and_fan_in_trunks_are_not_false_aliasing():
    doc = _doc([
        _node("src", 300, 60, style={"labelPosition": "none"}),
        _node("a", 150, 300, style={"labelPosition": "none"}),
        _node("b", 450, 300, style={"labelPosition": "none"}),
        _node("sink", 300, 540, style={"labelPosition": "none"}),
    ], [
        {"from": "src", "to": "a"}, {"from": "src", "to": "b"},
        {"from": "a", "to": "sink"}, {"from": "b", "to": "sink"},
    ])
    diagram, root, _ = _build(doc)
    assert link_aliasing_warnings(root, diagram.links) == []


def test_default_container_label_reserves_its_band():
    # A `vpc` with no `label` still draws its registry default "VPC": the band
    # is reserved so auto-placed children never sit under it.
    doc = _doc([{"kind": "container", "id": "v", "type": "vpc", "provider": "aws", "children": [_node("a")]}])
    diagram, root, _ = _build(doc)
    vpc = next(b for b in iter_boxes(root) if b.element.id == "v")
    assert vpc.label_text == "VPC" and vpc.label_reserve > 0
    assert diagram_warnings(diagram, root, REGISTRY) == []


def test_bottom_label_reserves_the_bottom_band():
    doc = _doc([{"kind": "container", "id": "g", "type": "group", "label": "Bottom",
                 "style": {"labelPosition": "bottom-left"}, "children": [_node("a")]}])
    diagram, root, _ = _build(doc)
    g = next(b for b in iter_boxes(root) if b.element.id == "g")
    a = next(b for b in iter_boxes(root) if b.element.id == "a")
    assert a.abs_y - g.abs_y == pytest.approx(32)  # just the padding at the top
    assert diagram_warnings(diagram, root, REGISTRY) == []


def test_right_label_is_part_of_the_footprint():
    doc = _doc([{"kind": "container", "id": "g", "type": "group", "layout": {"direction": "horizontal", "gap": 8},
                 "children": [_node("a", label="Primary database", style={"labelPosition": "right"}), _node("b")]}])
    diagram, root, _ = _build(doc)
    a = next(b for b in iter_boxes(root) if b.element.id == "a")
    assert a.footprint_w > a.width + 60
    assert diagram_warnings(diagram, root, REGISTRY) == []


def test_preview_warns_when_no_cjk_font_is_available(monkeypatch):
    import zook.preview

    monkeypatch.setattr(zook.preview, "_font_choice", lambda override: (None, None))
    diagram, root, _ = _build(_doc([_node("a", label="データベース")]))
    (message,) = zook.preview.preview_font_warnings(diagram, root, REGISTRY)
    assert "ZOOK_PREVIEW_FONT" in message


def test_an_unusable_preview_font_override_is_ignored_and_reported(monkeypatch, tmp_path):
    import zook.preview

    bogus = tmp_path / "typo.ttf"
    bogus.write_bytes(b"not a font")
    monkeypatch.setenv("ZOOK_PREVIEW_FONT", str(bogus))
    monkeypatch.setattr(zook.preview, "_FONT_CANDIDATES", [])
    zook.preview._font_choice.cache_clear()
    try:
        diagram, root, _ = _build(_doc([_node("a", label="Database")]))
        (message,) = zook.preview.preview_font_warnings(diagram, root, REGISTRY)
        assert "typo.ttf" in message and "ignored" in message
        zook.preview.render_preview(diagram, root, REGISTRY)  # still draws, with the fallback font
    finally:
        zook.preview._font_choice.cache_clear()


def test_a_self_loop_goes_round_a_corner_of_its_node():
    doc = _doc([_node("a", 200, 200, label="Job")], [{"from": "a", "to": "a", "label": "retry"}])
    diagram, root, _ = _build(doc)
    boxes = {b.element.id: b for b in iter_boxes(root)}
    a = boxes["a"]
    start, end, style, path = link_render_plan(a, a, diagram.links[0])
    assert (start, end, style, len(path)) == (3, 0, "polyline", 5)
    assert path[0] == (a.abs_x + a.width, a.abs_y + a.height / 2)
    assert path[-1] == (a.abs_x + a.width / 2, a.abs_y)
    assert diagram_warnings(diagram, root, REGISTRY) == []


def test_a_link_to_the_enclosing_container_meets_its_frame():
    doc = _doc([
        {"kind": "container", "id": "g", "type": "group", "label": "G", "layout": {"direction": "horizontal"},
         "children": [_node("a"), _node("b")]},
    ], [{"from": "g", "to": "a"}])
    diagram, root, _ = _build(doc)
    boxes = {b.element.id: b for b in iter_boxes(root)}
    _, _, style, path = link_render_plan(boxes["g"], boxes["a"], diagram.links[0])
    assert style == "straight" and len(path) == 2
    g = boxes["g"]
    (x, y) = path[0]
    assert x in (g.abs_x, g.abs_x + g.width) or y in (g.abs_y, g.abs_y + g.height)
    assert diagram_warnings(diagram, root, REGISTRY) == []


def test_pushing_past_an_explicit_sibling_does_not_land_on_another_auto_one():
    doc = _doc([
        {"kind": "container", "id": "g", "type": "group", "layout": {"direction": "grid", "columns": 2},
         "children": [_node("e", 40, 70), _node("a1"), _node("a2"), _node("a3"), _node("a4")]},
    ])
    diagram, root, _ = _build(doc)
    assert not [w for w in diagram_warnings(diagram, root, REGISTRY) if "overlaps element" in w]


def _cxn_ids(connector):
    from pptx.oxml.ns import qn

    nv = connector._element.find(qn("p:nvCxnSpPr")).find(qn("p:cNvCxnSpPr"))
    st, end = nv.find(qn("a:stCxn")), nv.find(qn("a:endCxn"))
    return (st.get("id") if st is not None else None), (end.get("id") if end is not None else None)


def test_a_connector_leaving_past_a_label_is_glued_to_that_label():
    # A viewer that re-snaps glued ends (LibreOffice on open, PowerPoint
    # after a move) must land where the route was planned: under the label,
    # not on the icon's edge with the line through the label text.
    doc = _doc([_node("a", 100, 100, label="Upper"), _node("b", 100, 400, label="Lower")], [{"from": "a", "to": "b"}])
    _, _, prs = _build(doc)
    shapes = list(_all_shapes(prs.slides[0].shapes))
    by_name = {s.name: s for s in shapes}
    (conn,) = [s for s in shapes if s._element.tag.endswith("cxnSp")]
    assert _cxn_ids(conn) == (str(by_name["a label"].shape_id), str(by_name["b icon"].shape_id))


def test_parallel_links_are_left_unglued_on_their_lanes():
    doc = _doc([_node("a", 100, 100), _node("b", 400, 100)], [{"from": "a", "to": "b"}, {"from": "b", "to": "a"}])
    _, _, prs = _build(doc)
    connectors = [s for s in _all_shapes(prs.slides[0].shapes) if s._element.tag.endswith("cxnSp")]
    assert len(connectors) == 2 and all(_cxn_ids(c) == (None, None) for c in connectors)


def test_a_node_is_one_group_named_by_its_id():
    doc = _doc([_node("web", 100, 100, label="Web server")])
    _, _, prs = _build(doc)
    (group,) = [s for s in prs.slides[0].shapes if s.name == "web"]
    assert {s.name for s in group.shapes} == {"web icon", "web label"}
    pic = next(s for s in group.shapes if s.name == "web icon")
    assert pic._element.nvPicPr.cNvPr.get("descr") == "Web server (EC2)"


def test_the_pptx_says_zook_made_it_and_what_size_it_is(monkeypatch):
    import datetime
    import re as _re

    monkeypatch.setenv("SOURCE_DATE_EPOCH", "1767225600")  # 2026-01-01
    _, _, prs = _build(_doc([_node("a", 100, 100)]))
    props = prs.core_properties
    assert props.author == props.last_modified_by == "zook"
    assert props.description.startswith("Generated by zook")
    assert props.created == datetime.datetime(2026, 1, 1)
    from pptx.oxml.ns import qn

    assert "type" not in prs.part._element.find(qn("p:sldSz")).attrib
    app = next(p for p in prs.part.package.iter_parts() if str(p.partname) == "/docProps/app.xml").blob.decode()
    assert "<PresentationFormat>Custom</PresentationFormat>" in app and _re.search(r"<Slides>1</Slides>", app)


def test_japanese_text_is_marked_as_japanese():
    _, _, prs = _build(_doc([_node("a", 100, 100, label="ウェブサーバー")]))
    label = next(s for s in _all_shapes(prs.slides[0].shapes) if s.name == "a label")
    run = label.text_frame.paragraphs[0].runs[0]
    assert run._r.rPr.get("lang") == "ja-JP"


def test_labels_and_lines_stay_readable_on_a_dark_canvas():
    doc = _doc([_node("a", 100, 100, label="Web"), _node("b", 400, 100, label="DB")], [{"from": "a", "to": "b"}],
               background="#1E2A38")
    _, _, prs = _build(doc)
    shapes = list(_all_shapes(prs.slides[0].shapes))
    label = next(s for s in shapes if s.name == "a label")
    assert str(label.text_frame.paragraphs[0].runs[0].font.color.rgb) == "FFFFFF"
    (conn,) = [s for s in shapes if s._element.tag.endswith("cxnSp")]
    assert str(conn.line.color.rgb) == "FFFFFF"


def test_a_word_in_a_box_sized_to_it_stays_on_one_line():
    from zook.layout import link_label_box

    for word in ("async", "sync", "metrics", "HTTPS", "データ"):
        assert link_label_box(word, 8)[2] == [word]


def test_link_color_dash_and_width_reach_the_pptx():
    from pptx.oxml.ns import qn

    doc = _doc([_node("a", 100, 100), _node("b", 400, 100), _node("c", 700, 100)], [
        {"from": "a", "to": "b", "color": "#E7157B", "line": "dashed", "width": 3},
        {"from": "b", "to": "c", "line": "dotted"},
    ])
    _, _, prs = _build(doc)
    first, second = [s for s in _all_shapes(prs.slides[0].shapes) if s._element.tag.endswith("cxnSp")]
    assert str(first.line.color.rgb) == "E7157B" and first.line.width == Pt(3)
    assert first.line._get_or_add_ln().find(qn("a:prstDash")).get("val") == "dash"
    assert second.line._get_or_add_ln().find(qn("a:prstDash")).get("val") == "sysDot"
    assert str(second.line.color.rgb) == "545B64" and second.line.width == Pt(1.25)


def test_an_invalid_link_colour_or_line_is_a_schema_error():
    from zook.errors import DiagramError

    for bad in ({"color": "pink"}, {"line": "wavy"}, {"width": 0}):
        with pytest.raises(DiagramError, match="links"):
            _doc([_node("a"), _node("b")], [{"from": "a", "to": "b", **bad}])
